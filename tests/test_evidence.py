import hashlib,json,tempfile,unittest
from pathlib import Path
from cryptography.hazmat.primitives import serialization
from cryptography.hazmat.primitives.asymmetric.ed25519 import Ed25519PrivateKey
from PIL import Image
import numpy as np
from verifier.provenance import canonical,verify_record,verify_captures,sha256_file,GENESIS
from verifier.statistics import recompute,compare_reported,DataError
from verifier.images import similarity,possible_reuse
from verifier.parser import parse_document,text_batches,ParseError
from verifier.data import data_checks
from verifier.local import analyse,InputError

class EvidenceTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
        self.key=Ed25519PrivateKey.generate()
        self.pub=self.key.public_key().public_bytes(serialization.Encoding.PEM,serialization.PublicFormat.SubjectPublicKeyInfo).decode()
        self.csv=self.root/'data.csv';self.csv.write_text('group,value,before,after\nA,1,1,2\nA,2,2,4\nA,3,3,4\nB,5,4,6\nB,6,5,8\nB,7,6,7\n')
        self.record=self.sign(self.csv)
        self.capture={'id':self.record['id'],'path':str(self.csv),'record':self.record,'org_id':'o','project_id':'p'}
        self.dataset={'id':'d','label':'Dataset 1','path':str(self.csv),'capture_id':self.capture['id'],'org_id':'o','project_id':'p','unit':'mg'}


    def tearDown(self):
        self.tmp.cleanup()


    def sign(self,path,prev=GENESIS,cid='cap_1'):
        record={'id':cid,'type':'reading','file_name':Path(path).name,'file_sha256':sha256_file(path),'captured_at':'2026-10-11T00:00:00+00:00','device_id':'device','location':'lab','prev_hash':prev}
        record['record_hash']=hashlib.sha256(canonical(record)).hexdigest()
        record['signature']=self.key.sign(bytes.fromhex(record['record_hash'])).hex()
        return record


    def provenance(self,captures=None,heads=None):
        return verify_captures(captures or [self.capture],{'device':self.pub},org_id='o',project_id='p',expected_heads=heads or {'device':self.record['record_hash']})


    def test_signature_and_byte_tamper(self):
        self.assertEqual(verify_record(self.record,self.csv,self.pub),'VERIFIED')
        self.csv.write_text('changed')
        self.assertEqual(verify_record(self.record,self.csv,self.pub),'TAMPERED_FILE')


    def test_record_tamper(self):
        self.assertEqual(verify_record(self.record|{'location':'elsewhere'},self.csv,self.pub),'TAMPERED_RECORD')


    def test_signature_wrong_key(self):
        other=Ed25519PrivateKey.generate().public_key().public_bytes(serialization.Encoding.PEM,serialization.PublicFormat.SubjectPublicKeyInfo)
        self.assertEqual(verify_record(self.record,self.csv,other),'BAD_SIGNATURE')


    def test_chain_integrity_and_anchor(self):
        self.assertTrue(self.provenance()['cap_1']['valid'])
        no_head=verify_captures([self.capture],{'device':self.pub},org_id='o',project_id='p')
        self.assertFalse(no_head['cap_1']['valid'])
        self.assertEqual(no_head['cap_1']['status'],'UNANCHORED_CHAIN')
        wrong=self.provenance(heads={'device':GENESIS})
        self.assertEqual(wrong['cap_1']['status'],'BROKEN_CHAIN')


    def test_chain_middle_deletion(self):
        second=self.sign(self.csv,self.record['record_hash'],'cap_2')
        third=self.sign(self.csv,second['record_hash'],'cap_3')
        third_cap=self.capture|{'id':'cap_3','record':third}
        self.assertFalse(self.provenance([self.capture,third_cap],{'device':third['record_hash']})['cap_3']['valid'])


    def test_tenant_isolation(self):
        cap=self.capture|{'org_id':'other'}
        result=verify_captures([cap],{'device':self.pub},org_id='o',project_id='p')
        self.assertEqual(result['cap_1']['status'],'SCOPE_MISMATCH')
        with self.assertRaises(InputError): analyse('',context={'org_id':'o','project_id':'p','captures':[cap]})


    def test_unknown_device(self):
        result=verify_captures([self.capture],{},org_id='o',project_id='p')
        self.assertEqual(result['cap_1']['status'],'UNKNOWN_DEVICE')


    def test_mean_n_and_std(self):
        result=recompute(self.csv,{'column':'value','metric':'mean','group_column':'group','group_a':'A'})
        self.assertEqual(result['mean'],2);self.assertEqual(result['n'],3);self.assertEqual(result['std'],1)
        self.assertTrue(compare_reported(2.00,result['mean'],metric='mean',decimals=2))
        self.assertFalse(compare_reported(99,result['mean'],metric='mean'))


    def test_welch_and_paired(self):
        result=recompute(self.csv,{'column':'value','metric':'p_value','group_column':'group','group_a':'A','group_b':'B','test':'welch'})
        self.assertLess(result['p_value'],.05)
        paired=recompute(self.csv,{'column':'before','paired_column':'after','metric':'p_value','test':'paired'})
        self.assertEqual(paired['paired_n'],6);self.assertLess(paired['p_value'],.05)


    def test_missing_paired_rows_do_not_shift(self):
        path=self.root/'pairs.csv';path.write_text('a,b\n1,2\n,4\n3,\n4,6\n')
        result=recompute(path,{'column':'a','paired_column':'b','metric':'p_value','test':'paired'})
        self.assertEqual(result['paired_n'],2);self.assertEqual(result['paired_missing'],2)


    def test_nonfinite_and_undefined_test(self):
        self.csv.write_text('value\nNaN\n')
        with self.assertRaises(DataError): recompute(self.csv,{'column':'value','metric':'mean'})
        with self.assertRaises(DataError): recompute(self.csv,{'column':'bad','metric':'n'})


    def test_numeric_mismatch_and_unsigned(self):
        claim={'claim':'Dataset 1 mean is 99 mg.','dataset_id':'d','column':'value','metric':'mean','reported':99}
        results=data_checks(claim['claim'],[claim],[self.dataset],[],[self.capture],self.provenance(),org_id='o',project_id='p')
        self.assertIn('DATA_MISMATCH',{r['status'] for r in results})
        results=data_checks(claim['claim'],[claim],[self.dataset],[],[self.capture],{},org_id='o',project_id='p')
        self.assertIn('UNSIGNED_DATA',{r['status'] for r in results})


    def test_image_hash_resize_crop_and_mismatch(self):
        rng=np.random.default_rng(5);pixels=rng.integers(0,256,(128,128,3),dtype=np.uint8)
        a=self.root/'a.png';b=self.root/'b.png';c=self.root/'c.png';d=self.root/'d.png'
        Image.fromarray(pixels).save(a)
        Image.open(a).resize((256,256)).save(b)
        Image.open(a).crop((32,32,96,96)).save(c)
        Image.new('RGB',(128,128),'white').save(d)
        self.assertEqual(similarity(a,a)['kind'],'exact_hash')
        self.assertEqual(similarity(b,a)['kind'],'resize_similar')
        self.assertEqual(similarity(c,a,crop=[.25,.25,.75,.75])['kind'],'crop_or_resize_similar')
        self.assertEqual(similarity(a,d)['kind'],'different')


    def test_perceptual_collision_not_verified(self):
        a=self.root/'dark.png';b=self.root/'light.png'
        Image.new('RGB',(32,32),'black').save(a);Image.new('RGB',(32,32),'white').save(b)
        self.assertEqual(similarity(a,b)['kind'],'different')


    def test_automatic_crop_proposal(self):
        import cv2
        pixels=np.random.default_rng(8).integers(0,256,(640,640,3),dtype=np.uint8)
        pixels=cv2.GaussianBlur(pixels,(5,5),0)
        original=self.root/'original.png';cropped=self.root/'cropped.png'
        Image.fromarray(pixels).save(original)
        Image.open(original).crop((160,160,480,480)).save(cropped)
        result=similarity(cropped,original)
        self.assertEqual(result['kind'],'automatic_crop_similar')
        self.assertIn('crop_proposal',result)


    def test_reuse_requires_distinct_conditions(self):
        a=self.root/'a.png';Image.new('RGB',(20,20),'green').save(a)
        figs=[{'id':'a','path':str(a),'condition':'control'},{'id':'b','path':str(a),'condition':'treated'}]
        self.assertEqual(len(possible_reuse(figs)),1)
        figs[1]['condition']='control';self.assertEqual(possible_reuse(figs),[])


    def test_parser_formats_and_boundary(self):
        md=self.root/'paper.md';md.write_text('# Results\nClaim.\n# References\n[1] Entry\n')
        parsed=parse_document(md);self.assertEqual(len(parsed['sections']),2)
        batches=list(text_batches(('Paragraph\n'*8000)+'\nReferences\n[1] Entry'))
        self.assertGreater(len(batches),1);self.assertTrue(all('[1] Entry' in b for b in batches))
        md.write_text('![image](../outside.png)')
        with self.assertRaises(ParseError): parse_document(md)


    def test_docx_parser(self):
        from docx import Document
        doc=Document();doc.add_paragraph('Measured result');doc.add_table(rows=2,cols=2)
        path=self.root/'paper.docx';doc.save(path)
        parsed=parse_document(path);self.assertIn('Measured result',parsed['text']);self.assertEqual(len(parsed['tables']),1)


    def test_overclaim_wording_is_review_flag(self):
        text='This treatment cures all disease.'
        findings=data_checks(text,[],[],[],[],{},org_id='o',project_id='p')
        self.assertEqual(findings[0]['status'],'OVERCLAIM');self.assertTrue(findings[0]['needs_human_review'])


    def test_complete_signed_data_report(self):
        text='Dataset 1 mean is 4 mg.'
        context={'org_id':'o','project_id':'p','captures':[self.capture],'datasets':[self.dataset],
            'public_keys':{'device':self.pub},'expected_heads':{'device':self.record['record_hash']},
            'claims':[{'claim':text,'dataset_id':'d','column':'value','metric':'mean','reported':4}]}
        report=analyse(text,context=context)
        self.assertTrue(report['provenance']['cap_1']['valid'])
        self.assertTrue(any(f.get('capture_id')=='cap_1' and f['status']=='VERIFIED' for f in report['flags']))
        self.assertTrue(any(f.get('claim')==text and f['status']=='VERIFIED' for f in report['flags']))
        self.csv.write_text('tampered')
        report=analyse(text,context=context)
        self.assertEqual(report['provenance']['cap_1']['status'],'TAMPERED_FILE')
        self.assertIn('DATA_MISMATCH',report['summary'])
