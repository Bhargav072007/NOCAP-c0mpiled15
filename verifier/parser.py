"""Local PDF/DOCX/TeX/Markdown parsing; retain page and asset locations."""
import hashlib
import re
from pathlib import Path

class ParseError(ValueError): pass

def safe_relative(base,name):
    path=(base/name).resolve()
    if not path.is_relative_to(base.resolve()): raise ParseError('Referenced file escapes document directory')
    return path

def split_bibliography(text):
    match=re.search(r'(?im)^\s*(?:#+\s*)?(?:references|bibliography|works cited)\s*:?[ \t]*$',text)
    return (text[:match.start()],text[match.start():]) if match else (text,'')

def sections(text):
    matches=list(re.finditer(r'(?m)^(?:#{1,6}\s+.+|(?:\d+(?:\.\d+)*[ .]+)?(?:Abstract|Introduction|Methods|Materials and Methods|Results|Discussion|Conclusion|References|Bibliography))\s*$',text))
    if not matches: return [{'id':'section_1','heading':'Document','start':0,'end':len(text),'text':text}]
    chunks=[]
    if matches[0].start()>0: chunks.append({'id':'section_0','heading':'Preamble','start':0,'end':matches[0].start(),'text':text[:matches[0].start()]})
    for i,m in enumerate(matches):
        end=matches[i+1].start() if i+1<len(matches) else len(text)
        chunks.append({'id':f'section_{i+1}','heading':m.group().strip(),'start':m.start(),'end':end,'text':text[m.start():end]})
    return chunks

def text_batches(text,limit=24000,*,include_bibliography=True):
    body,bib=split_bibliography(text)
    if len(bib)>25000: raise ParseError('Bibliography exceeds 25,000 characters; provide parsed references separately')
    available=min(limit,59000-len(bib))
    start=0
    while start<len(body):
        end=min(start+available,len(body))
        if end<len(body):
            boundary=body.rfind('\n',start,end)
            if boundary>start+available//2: end=boundary+1
        # overlap preserves claims crossing a chunk boundary; deduplication occurs after extraction.
        yield body[start:end]+('\n\n'+bib if include_bibliography else '')
        if end==len(body): break
        start=max(start+1,end-600)
    if not body.strip() and bib and include_bibliography: yield bib

def parse_document(path,asset_dir=None):
    path=Path(path).resolve();suffix=path.suffix.lower()
    if path.stat().st_size>50_000_000: raise ParseError('Document exceeds 50 MB')
    assets=Path(asset_dir) if asset_dir else None
    if assets: assets.mkdir(parents=True,exist_ok=True)
    pages=[];figures=[];tables=[];warnings=[]
    if suffix=='.pdf':
        import pdfplumber
        from pypdf import PdfReader
        reader=PdfReader(path)
        for i,page in enumerate(reader.pages):
            pages.append({'page':i+1,'text':page.extract_text() or ''})
            if assets:
                for j,image in enumerate(page.images):
                    dest=assets/f'page_{i+1}_image_{j+1}_{Path(image.name).name}'
                    dest.write_bytes(image.data)
                    figures.append({'id':f'pdf_p{i+1}_image{j+1}','path':str(dest.resolve()),'page':i+1,'label':'','origin':'embedded_pdf_asset'})
        with pdfplumber.open(path) as pdf:
            for i,page in enumerate(pdf.pages):
                for j,table in enumerate(page.extract_tables()):
                    tables.append({'id':f'table_p{i+1}_{j+1}','page':i+1,'rows':table})
        if not any(p['text'].strip() for p in pages): warnings.append('Scanned PDF needs OCR; no text extracted')
    elif suffix=='.docx':
        from docx import Document
        doc=Document(path);pages=[{'page':None,'text':'\n'.join(p.text for p in doc.paragraphs)}]
        tables=[{'id':f'table_{i+1}','rows':[[c.text for c in row.cells] for row in t.rows]} for i,t in enumerate(doc.tables)]
        if assets:
            for i,rel in enumerate(doc.part.rels.values()):
                if 'image' in rel.reltype and not rel.is_external:
                    dest=assets/f'docx_image_{i}{Path(rel.target_part.partname).suffix}'
                    dest.write_bytes(rel.target_part.blob)
                    figures.append({'id':f'docx_image_{i}','path':str(dest.resolve()),'label':'','origin':'embedded_docx_asset'})
    elif suffix=='.tex':
        text=path.read_text();text=re.sub(r'(?m)(?<!\\)%.*$','',text)
        for name in re.findall(r'\\bibliography\{([^}]+)\}',text):
            for entry in name.split(','):
                bib=safe_relative(path.parent,entry if entry.endswith('.bib') else entry+'.bib')
                if bib.exists(): text+='\nBibliography\n'+bib.read_text()
                else: warnings.append('Missing bibliography: '+entry)
        for name in re.findall(r'\\addbibresource(?:\[[^]]*\])?\{([^}]+)\}',text):
            bib=safe_relative(path.parent,name)
            if bib.exists(): text+='\nBibliography\n'+bib.read_text()
        for i,name in enumerate(re.findall(r'\\includegraphics(?:\[[^]]*\])?\{([^}]+)\}',text)):
            image=safe_relative(path.parent,name)
            if image.exists(): figures.append({'id':f'tex_image_{i+1}','path':str(image),'label':name,'origin':'tex_asset'})
        pages=[{'page':None,'text':text}]
    elif suffix in ('.txt','.md','.markdown'):
        text=path.read_text();pages=[{'page':None,'text':text}]
        for i,(label,name) in enumerate(re.findall(r'!\[([^]]*)\]\(([^)]+)\)',text)):
            if re.match(r'https?://',name): warnings.append('Remote manuscript image not fetched');continue
            image=safe_relative(path.parent,name)
            if image.exists(): figures.append({'id':f'md_image_{i+1}','path':str(image),'label':label,'origin':'markdown_asset'})
    else: raise ParseError('Accepted document formats: PDF, DOCX, TeX, Markdown, text')
    text='\n\n'.join(p['text'] for p in pages)
    if tables: text+='\n\nExtracted tables\n'+'\n'.join(str(t['rows']) for t in tables)
    return {'text':text,'sections':sections(text),'pages':pages,'figures':figures,'tables':tables,
            'file_sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'warnings':warnings}
